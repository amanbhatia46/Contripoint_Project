import unittest
import time
import HtmlTestRunner as x
import pyautogui
from multiprocessing import Event
from traceback import print_exc
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.support.select import Select
from webdriver_manager.chrome import ChromeDriverManager
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.common.action_chains import ActionChains
from selenium.webdriver.common.desired_capabilities import DesiredCapabilities
from openpyxl import load_workbook
from datetime import datetime, timedelta

from LoginContripoint import *


class Login(unittest.TestCase):

    def __init__(self):
        self.filename = '/home/am.bhatia/Desktop/contripoint/abc1.xlsx'
        self.wb = load_workbook(filename=self.filename)
        self.index_sheet = 0
        if 'Login' not in self.wb.sheetnames:
            self.wb.create_sheet('Login')
            self.wb.save('/home/am.bhatia/Desktop/contripoint/abc1.xlsx')
        self.wb = load_workbook(filename=self.filename)
        self.ws = self.wb.active
        for i, n in enumerate(self.wb.sheetnames):
            if n == 'Login':
                self.index_sheet = i
                break
        self.wb.active = self.index_sheet
        self.ws = self.wb.active
        self.wb.active = 1
        self.ws['A1'] = 'Email'
        self.ws['B1'] = 'Password'
        self.ws['C1'] = 'Test Case Module'
        self.ws['G1'] = 'Result'
        self.ws['H1'] = 'Date and Time'
        self.wb.save('/home/am.bhatia/Desktop/contripoint/abc1.xlsx')

        # datetime object containing current date and time
        futuredate = str(datetime.now())
        print(futuredate)
        self.ws['H2'] = futuredate

    def setExternalDriver(self, driver):
        self.driver = driver
        x = self.driver.window_handles[0]
        time.sleep(6)
        self.driver.switch_to.window(x)

    def tearDown(self):
        self.driver.quit()


if __name__ == '__main__':
    tb = Login()
    login = Login()
    driver = login.LoginMFA()
    tb.setExternalDriver(driver=driver)